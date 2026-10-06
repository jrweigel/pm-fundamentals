from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from common import load_yaml


def join_values(value: Any) -> str:
    return "; ".join(str(item) for item in value) if isinstance(value, list) else str(value or "")


def write_rows(sheet, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def main() -> None:
    parser = argparse.ArgumentParser(description="Render PM YAML trackers to XLSX")
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    plan = load_yaml(args.project / "plan.yaml")
    raid = load_yaml(args.project / "raid.yaml")
    project = load_yaml(args.project / "project.yaml")
    outcomes = project.get("outcomes", [])
    outcome_by_id: dict[str, dict[str, Any]] = {}
    for outcome in outcomes:
        outcome_id = str(outcome.get("outcomeId", ""))
        if not outcome_id:
            raise ValueError("Each project outcome needs an outcomeId")
        if outcome_id in outcome_by_id:
            raise ValueError(f"Duplicate outcomeId: {outcome_id}")
        if not outcome.get("outcomeOwner"):
            raise ValueError(f"{outcome_id} needs an outcomeOwner")
        outcome_by_id[outcome_id] = outcome

    for task in plan.get("tasks", []):
        task_id = task.get("taskId", "<missing taskId>")
        if not task.get("deliverableOwner"):
            raise ValueError(f"{task_id} needs a deliverableOwner")
        outcome_id = str(task.get("outcomeId", ""))
        if not outcome_id or outcome_id not in outcome_by_id:
            raise ValueError(f"{task_id} must link to a valid outcomeId in project.yaml")

    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    plan_sheet = workbook.create_sheet("Project Plan")
    plan_headers = [
        "Task ID", "Tasks/Milestones", "Deliverable Owner", "Outcome ID", "Outcome Owner",
        "Start Date", "End Date", "State", "% Complete", "Acceptance Criteria",
        "Evidence Source", "Confirmation Status", "Dependency Unblocked", "Decision ID",
        "Notes", "Last Updated",
    ]
    plan_rows = [
        [
            task.get("taskId", ""), task.get("tasksMilestones", ""),
            task.get("deliverableOwner", ""), task.get("outcomeId", ""),
            outcome_by_id[str(task.get("outcomeId", ""))].get("outcomeOwner", ""),
            task.get("startDate", ""), task.get("endDate", ""), task.get("state", ""),
            task.get("percentComplete", 0), join_values(task.get("acceptanceCriteria", [])),
            task.get("evidenceSource", ""), task.get("confirmationStatus", ""),
            task.get("dependencyUnblocked", ""), task.get("decisionId", ""),
            task.get("notes", ""), task.get("lastUpdated", ""),
        ]
        for task in plan.get("tasks", [])
    ]
    write_rows(plan_sheet, plan_headers, plan_rows)
    plan_sheet.conditional_formatting.add(
        "H2:H1000",
        CellIsRule(
            operator="equal",
            formula=['"Complete"'],
            fill=PatternFill(fill_type="solid", fgColor="C6EFCE"),
        ),
    )

    risk_sheet = workbook.create_sheet("Risks")
    write_rows(
        risk_sheet,
        ["ID", "Risk Description", "Probability", "Severity", "Score",
         "Action Plan (Mitigation/Contingency)", "Owner", "Status", "Date Identified",
         "Evidence Source", "Decision ID"],
        [[risk.get("id", ""), risk.get("riskDescription", ""), risk.get("probability", ""),
          risk.get("severity", ""), risk.get("score", ""), risk.get("actionPlan", ""),
          risk.get("owner", ""), risk.get("status", ""), risk.get("dateIdentified", ""),
          risk.get("evidenceSource", ""), risk.get("decisionId", "")]
         for risk in raid.get("risks", [])],
    )
    issue_sheet = workbook.create_sheet("Issues")
    write_rows(
        issue_sheet,
        ["ID", "State", "Priority", "Category", "Issue Name", "Issue Description",
         "Submitted By", "Submitted Date", "Owner", "Notes/Resolution", "Resolution Date",
         "Evidence Source", "Decision ID"],
        [[issue.get("id", ""), issue.get("state", ""), issue.get("priority", ""),
          issue.get("category", ""), issue.get("issueName", ""), issue.get("issueDescription", ""),
          issue.get("submittedBy", ""), issue.get("submittedDate", ""), issue.get("owner", ""),
          issue.get("notesResolution", ""), issue.get("resolutionDate", ""),
          issue.get("evidenceSource", ""), issue.get("decisionId", "")]
         for issue in raid.get("issues", [])],
    )
    change_sheet = workbook.create_sheet("Changes")
    write_rows(
        change_sheet,
        ["ID", "Change Description & Justification", "Change Impact", "Requestor",
         "Request Date", "Approver", "Approval Date", "Notes", "Status", "Evidence Source"],
        [[change.get("id", ""), change.get("changeDescriptionJustification", ""),
          change.get("changeImpact", ""), change.get("requestor", ""),
          change.get("requestDate", ""), change.get("approver", ""),
          change.get("approvalDate", ""), change.get("notes", ""), change.get("status", ""),
          change.get("evidenceSource", "")]
         for change in raid.get("changes", [])],
    )
    for sheet in workbook.worksheets:
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(column[0].value or "")) + 2, 12), 45)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)


if __name__ == "__main__":
    main()
