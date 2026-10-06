from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from common import project_values, render


def main() -> None:
    parser = argparse.ArgumentParser(description="Render a PM Markdown outline to PPTX")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--project", type=Path)
    args = parser.parse_args()
    values = project_values(args.project or args.input.parent) if (args.project or args.input.parent / "project.yaml").exists() else {}
    source = render(args.input.read_text(encoding="utf-8"), values)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    current = None
    body: list[str] = []
    for line in source.splitlines():
        if line.startswith("# ") or line.startswith("## "):
            if current is not None:
                _slide(presentation, current, body)
            current, body = line.lstrip("# ").strip(), []
        elif line.strip():
            body.append(line.strip())
    if current is not None:
        _slide(presentation, current, body)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(args.output)


def _slide(presentation: Presentation, title: str, body: list[str]) -> None:
    layout = presentation.slide_layouts[5]
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(body):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line.lstrip("-* ")
        paragraph.font.size = Pt(22)


if __name__ == "__main__":
    main()
