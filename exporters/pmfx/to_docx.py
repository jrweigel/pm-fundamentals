from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document

from common import project_values, render


def add_markdown(document: Document, text: str) -> None:
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("|"):
            document.add_paragraph(stripped.replace("|", " | ").strip(" |"))
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(stripped)


def main() -> None:
    parser = argparse.ArgumentParser(description="Render a PM Markdown artifact to DOCX")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--project", type=Path)
    args = parser.parse_args()
    values = project_values(args.project or args.input.parent) if (args.project or args.input.parent / "project.yaml").exists() else {}
    with args.input.open(encoding="utf-8") as stream:
        source = render(stream.read(), values)
    document = Document()
    add_markdown(document, source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)


if __name__ == "__main__":
    main()
